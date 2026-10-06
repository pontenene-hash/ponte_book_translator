"""Editable DOCX plus paired PDFs with no silently truncated translation."""
import io, json, zipfile
from xml.sax.saxutils import escape
import fitz
from PIL import Image as PILImage
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.colors import HexColor
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Image
from docx import Document
from docx.shared import Inches, Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from core import figure_bytes

FONT='PonteJapanese'

def register_font():
    if FONT not in pdfmetrics.getRegisteredFontNames():
        # PyMuPDF ships this CJK font. Embed it so iPhone PDF viewers do not need local fonts.
        pdfmetrics.registerFont(TTFont(FONT,io.BytesIO(fitz.Font('japan').buffer)))

def _paragraphs(text,style):
    return [Paragraph(escape(line) if line else '&#160;',style) for line in text.split('\n')]

def _jp_pdf(page,include_figures=True,font_size=10.5):
    register_font()
    out=io.BytesIO()
    body=ParagraphStyle('Body',fontName=FONT,fontSize=font_size,leading=font_size*1.65,wordWrap='CJK',spaceAfter=5)
    title=ParagraphStyle('Title',parent=body,fontSize=16,leading=23,spaceAfter=12,textColor=HexColor('#214d47'))
    small=ParagraphStyle('Small',parent=body,fontSize=8.5,leading=13,textColor=HexColor('#52645f'))
    story=[Paragraph(escape(page['title'] or '日本語訳'),title)]
    if include_figures:
        for f in sorted(page['figures'],key=lambda x:(x['box'][1],x['box'][0])):
            data=figure_bytes(page,f)
            with PILImage.open(io.BytesIO(data)) as im: w,h=im.size
            scale=min(470/w,220/h)
            story.append(Image(io.BytesIO(data),width=w*scale,height=h*scale))
            if f['caption']: story.extend(_paragraphs(f['caption'],small))
            story.append(Spacer(1,8))
    story.extend(_paragraphs(page['translation'] or '［未翻訳］',body))
    if page['notes']:
        story.append(Spacer(1,12)); story.extend(_paragraphs('要確認：'+page['notes'],small))
    def footer(c,d):
        c.setFont(FONT,8)
        label=page['label'][:60]
        if label: c.drawString(42,24,label)
    SimpleDocTemplate(out,pagesize=A4,leftMargin=44,rightMargin=44,topMargin=40,bottomMargin=44,
        title=page['title'] or '日本語訳',author='PONTE Book Translator').build(story,onFirstPage=footer,onLaterPages=footer)
    return out.getvalue()

def export_pdfs(project,include_figures=True,font_size=10.5):
    spread=fitz.open(); alternate=fitz.open(); jp=fitz.open()
    w,h=A4
    for page in project['pages']:
        with fitz.open(stream=_jp_pdf(page,include_figures,font_size),filetype='pdf') as translated:
            for i in range(len(translated)):
                # Repeat the source for continuations, preserving left/right pairing in every format.
                p=spread.new_page(width=w*2,height=h)
                p.insert_image(fitz.Rect(15,20,w-15,h-20),stream=page['image'],keep_proportion=True)
                p.draw_line((w,18),(w,h-18),color=(.82,.85,.83),width=.6)
                p.show_pdf_page(fitz.Rect(w,0,w*2,h),translated,i)
                left=alternate.new_page(width=w,height=h)
                left.insert_image(fitz.Rect(15,20,w-15,h-20),stream=page['image'],keep_proportion=True)
                right=alternate.new_page(width=w,height=h)
                right.show_pdf_page(right.rect,translated,i)
                jp.insert_pdf(translated,from_page=i,to_page=i)
    if not len(spread): raise ValueError('出力するページがありません。')
    # First original is the left-hand page in two-page capable viewers.
    alternate.xref_set_key(alternate.pdf_catalog(),'PageLayout','/TwoPageLeft')
    results={}
    for name,doc in [('facing_pages.pdf',spread),('alternating_pages.pdf',alternate),('japanese_only.pdf',jp)]:
        doc.subset_fonts()
        results[name]=doc.tobytes(garbage=4,deflate=True); doc.close()
    return results

def export_docx(project,include_figures=True):
    doc=Document(); section=doc.sections[0]
    section.page_width=Inches(8.27); section.page_height=Inches(11.69)
    section.left_margin=section.right_margin=Inches(.6)
    style=doc.styles['Normal']; style.font.name='Yu Gothic'; style.font.size=Pt(10)
    style.element.rPr.rFonts.set(qn('w:eastAsia'),'游ゴシック')
    doc.core_properties.title=project['title']
    for i,page in enumerate(project['pages']):
        if i: doc.add_page_break()
        doc.add_heading(f'原書 {page["label"]}',level=1)
        with PILImage.open(io.BytesIO(page['image'])) as im: iw,ih=im.size
        width=min(6.8,8.9*iw/ih)
        doc.add_picture(io.BytesIO(page['image']),width=Inches(width))
        doc.add_page_break()
        doc.add_heading(page['title'] or '日本語訳',level=1)
        doc.add_paragraph(f'原書 {page["label"]} 対応・{page["status"]}')
        if include_figures:
            for f in sorted(page['figures'],key=lambda x:(x['box'][1],x['box'][0])):
                data=figure_bytes(page,f)
                with PILImage.open(io.BytesIO(data)) as im: fw,fh=im.size
                doc.add_picture(io.BytesIO(data),width=Inches(min(6.5,3.0*fw/fh)))
                if f['caption']: doc.add_paragraph(f['caption'])
        for paragraph in (page['translation'] or '［未翻訳］').split('\n'): doc.add_paragraph(paragraph)
        if page['notes']: doc.add_paragraph('要確認：'+page['notes'])
    out=io.BytesIO();doc.save(out);return out.getvalue()

def export_bundle(project,include_figures=True,font_size=10.5):
    files=export_pdfs(project,include_figures,font_size)
    files['translation.docx']=export_docx(project,include_figures)
    files['translation.txt']=('\n\n'.join(f'【原書 {p["label"]}】{p["title"]}\n{p["translation"]}\n要確認：{p["notes"]}' for p in project['pages'])).encode('utf-8-sig')
    out=io.BytesIO()
    with zipfile.ZipFile(out,'w',zipfile.ZIP_DEFLATED) as z:
        for n,b in files.items(): z.writestr(n,b)
    return files,out.getvalue()
